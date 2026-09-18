"""维修记录提取工具测试。"""

from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment

from func.extraction import extract_sheet_records
from func.maintenance_utils import parse_comment


def test_parse_comment_normalizes_day_and_night_shifts():
    assert parse_comment("白班：更换滤芯\n夜班: 检查液压油") == [
        ("day", "更换滤芯"),
        ("night", "检查液压油"),
    ]


def test_parse_comment_keeps_unmarked_shift():
    assert parse_comment("检查设备") == [("未标注", "检查设备")]


def test_extract_sheet_splits_long_unmarked_maintenance_across_shifts():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["设备", "原因", 1, 2, 3])
    worksheet.append(["TR001", "检修", 1440, 720, 1440])
    worksheet.cell(row=2, column=3).comment = Comment("长时间维修", "test")
    worksheet.cell(row=2, column=5).comment = Comment("白班：已标记维修", "test")

    records = extract_sheet_records(worksheet, 2026, 8)

    assert [(record["班次"], record["工时_分钟"]) for record in records] == [
        ("day", 720.0),
        ("night", 720.0),
        ("未标注", 720),
        ("day", 1440),
    ]


def test_maintenance_report_normalizes_chinese_shift_labels(monkeypatch, tmp_path):
    from func import excel_maintenance

    source = tmp_path / "input.xlsx"
    source.touch()
    records = [
        {
            "日期": date(2026, 8, 31),
            "原始设备名称": "TR001",
            "原因": "检修",
            "班次": "白班",
            "维修内容": "发动机异响",
            "工时_分钟": 30,
        },
        {
            "日期": date(2026, 8, 31),
            "原始设备名称": "TR001",
            "原因": "检修",
            "班次": "夜班",
            "维修内容": "液压油漏油",
            "工时_分钟": 30,
        },
    ]
    monkeypatch.setattr(excel_maintenance, "extract_all_records", lambda *args, **kwargs: records)

    sheets = excel_maintenance.process_maintenance_data(
        str(source),
        return_sheets=True,
        details_only=True,
    )

    assert sheets["维修明细"]["班次"].tolist() == ["day", "night"]


def test_maintenance_report_keeps_split_shifts_during_dedup(monkeypatch, tmp_path):
    from func import excel_maintenance

    source = tmp_path / "input.xlsx"
    source.touch()
    records = [
        {
            "日期": date(2026, 8, 31),
            "原始设备名称": "TR001",
            "原因": "检修",
            "班次": "day",
            "维修内容": "长时间维修",
            "工时_分钟": 720,
        },
        {
            "日期": date(2026, 8, 31),
            "原始设备名称": "TR001",
            "原因": "检修",
            "班次": "night",
            "维修内容": "长时间维修",
            "工时_分钟": 720,
        },
    ]
    monkeypatch.setattr(excel_maintenance, "extract_all_records", lambda *args, **kwargs: records)

    sheets = excel_maintenance.process_maintenance_data(
        str(source),
        return_sheets=True,
        details_only=True,
        use_ml_fallback=False,
    )

    assert sheets["维修明细"]["班次"].tolist() == ["day", "night"]
